"""
Compile manuscript to docx: pandoc for text+tables, python-docx for images.
Figures moved to first citation in text.
"""
import re, subprocess, os

ROOT = "/Users/cary/cea-intismeran"
os.chdir(ROOT)

# 1. Read manuscript
with open("manuscript.md") as f:
    text = f.read()
text = re.sub(r"^# Cost-Effectiveness[^\n]*\n", "", text)

# 2. Extract figures (caption + path), replace with nothing in-place
figures = {}  # fnum -> (caption, path)
def extract_figure(m):
    num = int(m.group(1)); caption = m.group(2).strip(); path = m.group(3)
    figures[num] = (caption, path)
    return ""  # remove from Figures section

text = re.sub(r'-\s*\*\*Figure\s+(\d+)[:\*]?\s*(.*?)\s*→\s*`([^`]+)`', extract_figure, text)
text = re.sub(r"\n## Figures\n", "\n", text)
text = re.sub(r"\n## Tables\n", "\n", text)

# 3. Tables: extract, embed at first citation
tables = {}
table_pat = re.compile(r'### (Table \d+\..*?)\n(.*?)(?=\n### |\n---\n|\Z)', re.DOTALL)
for m in table_pat.finditer(text):
    title = m.group(1); content = m.group(2)
    tnum = int(re.search(r'Table (\d+)', title).group(1))
    tables[tnum] = f"**{title}**\n\n{content.strip()}"
for tnum in sorted(tables.keys(), reverse=True):
    text = re.sub(rf'### Table {tnum}\..*?(?=\n### |\n---\n|\Z)', '', text, count=1, flags=re.DOTALL)
for tnum in sorted(tables.keys()):
    m = re.search(rf'\bTable\s+{tnum}\b', text)
    if m:
        rest = text[m.end():]
        eop = re.search(r'\n\n', rest)
        insert_at = m.end() + (eop.end() if eop else 0)
        text = text[:insert_at] + "\n\n" + tables[tnum] + "\n\n" + text[insert_at:]

# 4. Insert figure captions + image markers at first citation
for fnum in sorted(figures.keys(), reverse=True):  # reverse to preserve positions
    caption, _ = figures[fnum]
    m = re.search(rf'\bFigure\s+{fnum}\b', text)
    if m:
        rest = text[m.end():]
        eop = re.search(r'\n\n', rest)
        insert_at = m.end() + (eop.end() if eop else 0)
        block = f"\n\n**Figure {fnum}:** {caption}\n\n[IMG_{fnum}]\n\n"
        text = text[:insert_at] + block + text[insert_at:]

# 5. Cleanup + page breaks
text = re.sub(r'\n---\n', '\n\n', text); text = re.sub(r'\n{3,}', '\n\n', text)
text = re.sub(r'\n\n## Abstract\n', '\n\n\\newpage\n\n## Abstract\n', text, count=1)
text = re.sub(r'\nKeywords:[^\n]+\n\n', '\n\\newpage\n\n', text, count=1)

os.makedirs("build", exist_ok=True)
with open("build/title_block.md") as f:
    title_block = f.read()
with open("build/full_tmp.md", "w") as f:
    f.write(title_block + "\n\n" + text)

# 6. pandoc
subprocess.run(["pandoc", "build/full_tmp.md", "--reference-doc=manuscript_template.docx",
    "-f", "markdown+raw_tex", "-o", "output/manuscript_raw.docx"], check=True)

# 7. Insert images at [IMG_N] markers + page breaks
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document("output/manuscript_raw.docx")

def add_page_break_before(paragraph):
    run = paragraph.insert_paragraph_before().add_run()
    br = OxmlElement('w:br'); br.set(qn('w:type'), 'page')
    run._r.append(br)

for fnum in sorted(figures.keys()):
    _, img_path = figures[fnum]
    marker = f"[IMG_{fnum}]"
    for p in doc.paragraphs:
        if marker in p.text:
            for run in p.runs:
                if marker in run.text:
                    run.text = run.text.replace(marker, '')
            run = p.add_run()
            run.add_picture(img_path, width=Inches(5.5))
            print(f"Figure {fnum} inserted")
            break

# Page breaks: before Abstract, before 1. Introduction
paras = doc.paragraphs
in_abstract = False
for i, p in enumerate(paras):
    t = p.text.strip()
    s = p.style.name
    if s.startswith('Heading') and t == 'Abstract':
        add_page_break_before(p)
        print(f"break before Abstract (para {i})")
        in_abstract = True
    if s.startswith('Heading') and t.startswith('1. Introduction'):
        add_page_break_before(p)
        print(f"break before 1. Introduction (para {i})")
        in_abstract = False
    # Set Abstract paragraphs to single spacing to fit 1 page
    if in_abstract and not s.startswith('Heading'):
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)

doc.save("output/manuscript.docx")
print(f"Done -> output/manuscript.docx ({len(figures)} figs, {len(tables)} tables)")