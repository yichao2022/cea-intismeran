"""Compile manuscript.md into PharmacoEconomics submission-ready Word doc.
Tables and figures embedded in text at first citation.
"""
import re
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()

for section in doc.sections:
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)

style = doc.styles['Normal']
font = style.font; font.name = 'Times New Roman'; font.size = Pt(12)
style.paragraph_format.line_spacing = 2.0

def add_heading_styled(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs: run.font.color.rgb = RGBColor(0, 0, 0)

def add_para(text, bold=False, italic=False, indent=True):
    p = doc.add_paragraph()
    if indent: p.paragraph_format.first_line_indent = Cm(1.27)
    run = p.add_run(text); run.font.name = 'Times New Roman'; run.font.size = Pt(12)
    run.bold = bold; run.italic = italic
    p.paragraph_format.line_spacing = 2.0

def add_para_mixed(parts, indent=True):
    p = doc.add_paragraph()
    if indent: p.paragraph_format.first_line_indent = Cm(1.27)
    p.paragraph_format.line_spacing = 2.0
    for text, bold, italic in parts:
        run = p.add_run(text); run.font.name = 'Times New Roman'; run.font.size = Pt(12)
        run.bold = bold; run.italic = italic

def parse_para(para_text, indent=True):
    parts = []; idx = 0
    while idx < len(para_text):
        m = re.search(r'\*\*\*(.*?)\*\*\*', para_text[idx:])
        if m and m.start() == 0: parts.append((m.group(1), True, True)); idx += m.end(); continue
        m = re.search(r'\*\*(.*?)\*\*', para_text[idx:])
        if m and m.start() == 0: parts.append((m.group(1), True, False)); idx += m.end(); continue
        m = re.search(r'\*(.*?)\*', para_text[idx:])
        if m and m.start() == 0: parts.append((m.group(1), False, True)); idx += m.end(); continue
        m = re.search(r'^[^*\n]+', para_text[idx:])
        if m: parts.append((m.group(0), False, False)); idx += m.end()
        else: idx += 1
    if parts: add_para_mixed(parts, indent=indent)
    else: add_para(para_text, indent=indent)

def add_table_from_md(table_text):
    """Parse markdown table into docx table."""
    rows = [r.strip() for r in table_text.split('\n') if r.strip()]
    data_rows = [r for r in rows if not re.match(r'[\|\s\-:]+$', r)]
    if not data_rows: return
    headers = [h.strip() for h in data_rows[0].split('|') if h.strip()]
    ncols = len(headers); nrows = len(data_rows) - 1
    table = doc.add_table(rows=nrows + 1, cols=ncols)
    table.style = 'Table Grid'; table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]; cell.text = ''
        run = cell.paragraphs[0].add_run(h)
        run.bold = True; run.font.name = 'Times New Roman'; run.font.size = Pt(10)
    for row_idx, row_text in enumerate(data_rows[1:]):
        cells = [c.strip() for c in row_text.split('|') if c.strip()]
        for col_idx, cell_text in enumerate(cells):
            if col_idx < ncols:
                cell = table.rows[row_idx].cells[col_idx]; cell.text = ''
                ct = cell_text.strip()
                if ct.startswith('**') and ct.endswith('**'):
                    ct = ct[2:-2]; run = cell.paragraphs[0].add_run(ct); run.bold = True
                else: run = cell.paragraphs[0].add_run(ct)
                run.font.name = 'Times New Roman'; run.font.size = Pt(10)

# ── Load source files ──
with open("title_page.md") as f: title_text = f.read()
with open("manuscript.md") as f: text = f.read()

# ========== TITLE PAGE ==========
t = title_text.split("\n\n")[1].replace("## ", "").strip()
add_para("", indent=False); add_para("", indent=False)
add_para(t, bold=True, indent=False); add_para("", indent=False)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Yichao Jin, Ph.D.").font.name = 'Times New Roman'
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("School of Economic, Political and Policy Sciences, University of Texas at Dallas, Richardson, TX, USA").font.name = 'Times New Roman'
add_para("", indent=False)
add_para("Corresponding author: Yichao Jin — Yichao.Jin@UTDallas.edu", bold=True, indent=False)
add_para("ORCID: 0009-0003-7667-5143", indent=False); add_para("", indent=False)
add_para("Funding: This study was conducted without external funding.", indent=False)
add_para("Acknowledgments: None.", indent=False)
add_para("Conflict of interest: The author declares no competing interests relevant to this study.", indent=False)
add_para("Data and code availability: All data used in this analysis are derived from published sources cited in the references. The model code is available at https://github.com/yichao2022/cea-intismeran.", indent=False)
doc.add_page_break()

# ========== MANUSCRIPT BODY ==========
# Remove title line, strip Figures section (will embed)
text = re.sub(r"^# Cost-Effectiveness[^\n]*\n", "", text)
figures_section = ""
if "## Figures" in text:
    text, figures_section = text.split("## Figures", 1)

# Build figure lookup: fig_caption -> image_path
figures = {}
for line in figures_section.strip().split('\n'):
    m = re.search(r'-\s+\*\*Figure\s+(\d+)[:\*]?\s*(.*?)\s*→\s*`([^`]+)`', line)
    if m:
        fig_num = int(m.group(1))
        caption = m.group(2).strip()
        img_path = m.group(3)
        figures[fig_num] = {'caption': caption, 'path': img_path}

# Parse sections
sections = re.split(r'\n## ', text)

# Track which tables/figures have been placed
placed_tables = set()
placed_figures = set()

# Table lookup: extract all tables from manuscript
all_tables = {}  # table_num -> (title, md_table)
table_section_start = None
table_section_end = None

for i, sec in enumerate(sections):
    if sec.strip().startswith("Tables"):
        table_section_start = i
        break

if table_section_start is not None:
    tables_content = sections[table_section_start]
    # Remove the "Tables" heading line
    tables_content = re.sub(r'^Tables\n', '', tables_content)
    # Split on ### Table N
    table_blocks = re.split(r'\n### ', tables_content)
    for block in table_blocks:
        block = block.strip()
        if not block: continue
        lines = block.split('\n')
        title = lines[0].strip()
        # Get table number from title
        m = re.match(r'Table\s+(\d+)', title)
        if m:
            tnum = int(m.group(1))
            rest = '\n'.join(lines[1:]).strip()
            all_tables[tnum] = (title, rest)
    # Remove tables section from sections list
    del sections[table_section_start]

# Now process sections, embedding tables and figures at first mention
for section_text in sections:
    section_text = section_text.strip()
    if not section_text: continue

    lines = section_text.split('\n')
    section_title = lines[0].strip()
    content = '\n'.join(lines[1:]).strip()

    if section_title == "Abstract":
        doc.add_page_break()
        add_heading_styled("Abstract", level=1)
        for line in content.split('\n'):
            line = line.strip()
            if not line: continue
            m = re.match(r'\*\*(.*?):\*\*\s*(.*)', line)
            if m: parse_para(f"**{m.group(1)}:** {m.group(2)}")
            else: parse_para(line)

    elif section_title in ["1. Introduction", "2. Methods", "3. Results", "4. Discussion", "5. Declarations"]:
        num = section_title[0]
        clean_title = re.sub(r'^\d+\.\s*', '', section_title)
        add_heading_styled(f"{num}. {clean_title}", level=1)

        subsections = re.split(r'\n### ', content)
        for sub in subsections:
            sub = sub.strip()
            if not sub: continue
            sub_lines = sub.split('\n')
            sub_title = sub_lines[0].strip()
            sub_content = '\n'.join(sub_lines[1:]).strip()
            if sub_title == "---": continue
            add_heading_styled(sub_title, level=2)

            paragraphs = re.split(r'\n\n+', sub_content)
            for para in paragraphs:
                para = para.strip()
                if not para or para == "---": continue

                # Check for (Table X) or (Figure X) reference and embed
                for tnum in sorted(all_tables.keys()):
                    if tnum not in placed_tables and f"(Table {tnum})" in para:
                        # Add the table after this paragraph
                        pass  # Will handle below

                # Bullet list
                if para.startswith('- ') or para.startswith('* '):
                    for line in para.split('\n'):
                        line = line.strip()
                        if line.startswith('- '): line = line[2:]
                        elif line.startswith('* '): line = line[2:]
                        p = doc.add_paragraph(style='List Bullet')
                        run = p.add_run(line); run.font.name = 'Times New Roman'; run.font.size = Pt(12)
                        p.paragraph_format.line_spacing = 2.0
                    continue

                parse_para(para)

                # After paragraph, check if it referenced a table or figure
                for tnum in sorted(all_tables.keys()):
                    if tnum not in placed_tables:
                        if re.search(rf'\bTable\s+{tnum}\b', para):
                            title, md = all_tables[tnum]
                            add_para("", indent=False)
                            parse_para(f"**{title}**", indent=False)
                            add_table_from_md(md)
                            add_para("", indent=False)
                            placed_tables.add(tnum)

                for fnum in sorted(figures.keys()):
                    if fnum not in placed_figures:
                        if re.search(rf'\bFigure\s+{fnum}\b', para):
                            f = figures[fnum]
                            add_para("", indent=False)
                            parse_para(f"**Figure {fnum}:** {f['caption']}", indent=False)
                            try:
                                doc.add_picture(f['path'], width=Inches(5.5))
                            except:
                                pass
                            add_para("", indent=False)
                            placed_figures.add(fnum)

    elif section_title == "References":
        add_heading_styled("References", level=1)
        for line in content.split('\n'):
            line = line.strip()
            if line:
                p = doc.add_paragraph()
                run = p.add_run(line); run.font.name = 'Times New Roman'; run.font.size = Pt(12)
                p.paragraph_format.line_spacing = 2.0
                p.paragraph_format.left_indent = Cm(1.27)
                p.paragraph_format.first_line_indent = Cm(-1.27)

# ── Append any remaining tables/figures not referenced in text ──
if len(placed_tables) < len(all_tables):
    doc.add_page_break()
    add_heading_styled("Tables", level=1)
    for tnum in sorted(all_tables.keys()):
        if tnum not in placed_tables:
            title, md = all_tables[tnum]
            parse_para(f"**{title}**", indent=False)
            add_table_from_md(md)
            doc.add_paragraph()

if len(placed_figures) < len(figures):
    doc.add_page_break()
    add_heading_styled("Figures", level=1)
    for fnum in sorted(figures.keys()):
        if fnum not in placed_figures:
            f = figures[fnum]
            parse_para(f"**Figure {fnum}:** {f['caption']}", indent=False)
            try:
                doc.add_picture(f['path'], width=Inches(5.5))
            except:
                pass
            doc.add_paragraph()

output_path = "output/manuscript.docx"
doc.save(output_path)
print(f"✅ Saved to {output_path}")
print(f"   Tables embedded: {len(placed_tables)}/{len(all_tables)}, Figures embedded: {len(placed_figures)}/{len(figures)}")